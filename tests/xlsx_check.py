"""Opens an .xlsx made by web/js/xlsx.js with openpyxl, as Excel would: run by CI."""
import sys
import openpyxl

wb = openpyxl.load_workbook(sys.argv[1])
ws = wb["Customers"]
rows = [[c.value for c in r] for r in ws.iter_rows()]
assert rows[0][0] == "Name", rows
assert rows[1] == ["Ah Seng 阿成", 12, 123.45], rows
assert ws["A1"].font.b, "header not bold"
assert wb.sheetnames == ["Customers", "Second"], wb.sheetnames
print("xlsx opens:", rows)
